from __future__ import annotations

import csv
import io
import json
import os
import zipfile
from collections import Counter, OrderedDict
from contextlib import ExitStack
from pathlib import Path

from .models import BatchError, BatchSettings
from .storage import digest, file_manifest, read_json, resolve_file, write_json
from .tables import EXCEL_CELL_CHARS, INPUT_HEADER_PREFIX


RESULT_COLUMNS = ["pair_id", "a_row", "b_row", "a_id", "b_id", "a_name", "b_name",
                  "a_input_smiles", "b_input_smiles", "a_canonical_smiles", "b_canonical_smiles",
                  "candidate_index", "polymer_smiles", "polymer_class", "reaction_id",
                  "engine_mon1_smiles", "engine_mon2_smiles", "reactset"]
PAIR_COLUMNS = ["pair_id", "a_row", "b_row", "a_id", "b_id", "status", "candidate_count", "error_code", "message"]
ERROR_COLUMNS = ["role", "source_key", "row_number", "id", "input_smiles", "error_code", "message"]
EXCEL_ROWS = 1_000_000


def spreadsheet_text(value):
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


class OutputBudget:
    def __init__(self, limit: int):
        self.limit, self.used = limit, 0

    def add(self, size: int):
        self.used += size
        if self.used > self.limit:
            raise BatchError("结果数据超过存储限额，任务未完整导出。", "result_limit", 413)


class TableWriter:
    def __init__(self, directory: Path, name: str, columns: list[str], workbook, budget: OutputBudget, stack: ExitStack):
        self.name, self.columns, self.workbook, self.budget = name, columns, workbook, budget
        self.handle = stack.enter_context((directory / f"{name}.csv").open("w", encoding="utf-8-sig", newline=""))
        self.buffer = io.StringIO(newline="")
        self.writer = csv.writer(self.buffer)
        self.sheet_no, self.rows = 0, 0
        self._new_sheet()
        self._csv(columns)

    def _new_sheet(self):
        from openpyxl.cell import WriteOnlyCell
        if any(len(column) > EXCEL_CELL_CHARS for column in self.columns):
            raise BatchError("结果表头超过 Excel 可表示长度，未截断数据。", "excel_cell_limit", 413)
        self.sheet_no += 1
        self.sheet = self.workbook.create_sheet(self.name if self.sheet_no == 1 else f"{self.name}_{self.sheet_no}")
        self.sheet.freeze_panes = "A2"
        self.sheet.append([WriteOnlyCell(self.sheet, value=column) for column in self.columns])
        self.rows = 1

    def _csv(self, values):
        self.buffer.seek(0)
        self.buffer.truncate(0)
        self.writer.writerow(values)
        text = self.buffer.getvalue()
        self.budget.add(len(text.encode("utf-8")))
        self.handle.write(text)

    def append(self, record: dict):
        from openpyxl.cell import WriteOnlyCell
        values = [spreadsheet_text(record.get(column, "")) for column in self.columns]
        if any(isinstance(value, str) and len(value) > EXCEL_CELL_CHARS for value in values):
            raise BatchError("结果单元格超过 Excel 可表示长度，未截断数据。", "excel_cell_limit", 413)
        self._csv(values)
        if self.rows >= EXCEL_ROWS:
            self._new_sheet()
        cells = []
        for value in values:
            cell = WriteOnlyCell(self.sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
                cell.number_format = "@"
            cells.append(cell)
        self.sheet.append(cells)
        self.rows += 1


class PairIndex:
    """Index compact chemical pairs, never materialize expanded result rows."""
    def __init__(self, root: Path, chunks: list[dict], budget: int):
        self.locations: dict[tuple[str, str], tuple[Path, int]] = {}
        self.handles: OrderedDict = OrderedDict()
        size = 0
        for chunk in chunks:
            if chunk["kind"] != "generate":
                continue
            path = resolve_file(root, chunk["artifact"]["path"])
            size += path.stat().st_size
            if size > budget:
                raise BatchError("分片结果超过存储限额。", "result_limit", 413)
            if digest(path) != chunk["artifact"]["sha256"]:
                raise BatchError("分片校验失败。", "artifact_corrupt", 503)
            with path.open("rb") as handle:
                while True:
                    offset = handle.tell()
                    line = handle.readline()
                    if not line:
                        break
                    pair = json.loads(line)
                    key = pair["a"], pair["b"]
                    if key in self.locations:
                        raise RuntimeError("duplicate committed chemical pair")
                    self.locations[key] = path, offset

    def get(self, first: str, second: str) -> dict | None:
        location = self.locations.get((first, second))
        if location is None:
            return None
        path, offset = location
        handle = self.handles.pop(path, None)
        if handle is None:
            handle = path.open("rb")
        self.handles[path] = handle
        if len(self.handles) > 16:
            _, old = self.handles.popitem(last=False)
            old.close()
        handle.seek(offset)
        return json.loads(handle.readline())

    def close(self):
        for handle in self.handles.values():
            handle.close()


def export_job(root: Path, job: dict, snapshot: dict, chunks: list[dict], directory: Path, config: BatchSettings) -> dict:
    from openpyxl import Workbook
    directory.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    budget = OutputBudget(config.result_bytes)
    index = PairIndex(root, chunks, config.result_bytes)
    classification_errors = {}
    for chunk in chunks:
        if chunk["kind"] == "classify":
            artifact_path = resolve_file(root, chunk["artifact"]["path"])
            if digest(artifact_path) != chunk["artifact"]["sha256"]:
                raise BatchError("分类缓存校验失败。", "artifact_corrupt", 503)
            classification_errors.update(read_json(artifact_path).get("errors", {}))
    statuses: Counter = Counter()
    candidate_count = 0
    preview_offsets = []
    try:
        with ExitStack() as stack:
            tables = snapshot["tables"]
            results = TableWriter(directory, "results", RESULT_COLUMNS, workbook, budget, stack)
            pairs = TableWriter(directory, "pairs", PAIR_COLUMNS, workbook, budget, stack)
            errors = TableWriter(directory, "input_errors", ERROR_COLUMNS, workbook, budget, stack)
            preview_file = stack.enter_context((directory / "preview.jsonl").open("wb"))
            for role in ("a", "b"):
                columns = ["source_key", "row_number", "canonical_smiles", "error_code", "message"]
                # Prefix original headers so they cannot overwrite provenance.
                columns += [f"{INPUT_HEADER_PREFIX}{header}" for header in tables[role]["headers"]]
                inputs = TableWriter(directory, f"inputs_{role}", columns, workbook, budget, stack)
                for row in tables[role]["rows"]:
                    record = {key: row.get(key, "") for key in columns}
                    record.update({f"{INPUT_HEADER_PREFIX}{key}": value for key, value in row["values"].items()})
                    classification_error = classification_errors.get(row["canonical_smiles"])
                    if classification_error:
                        record.update(error_code="classification_error", message=classification_error)
                    inputs.append(record)
                    if row["error_code"] or classification_error:
                        errors.append({**row, "role": role, "error_code": record["error_code"], "message": record["message"]})
            for a in tables["a"]["rows"]:
                for b in tables["b"]["rows"]:
                    pair_id = f"a{a['row_number']}:b{b['row_number']}"
                    base = {"pair_id": pair_id, "a_row": a["row_number"], "b_row": b["row_number"], "a_id": a["id"], "b_id": b["id"]}
                    if not a["canonical_smiles"] or not b["canonical_smiles"]:
                        pair = {"status": "invalid_input", "error_code": "invalid_input", "message": "关联输入未通过预检。", "candidates": []}
                    elif classification_error := (classification_errors.get(a["canonical_smiles"]) or classification_errors.get(b["canonical_smiles"])):
                        # Cancellation can stop the job before generation starts.
                        # Already committed classification failures are known
                        # errors, even when that pair has no generation shard.
                        pair = {"status": "error", "error_code": "classification_error", "message": classification_error, "candidates": []}
                    else:
                        pair = index.get(a["canonical_smiles"], b["canonical_smiles"]) or {
                            "status": "not_processed", "error_code": "not_processed", "message": "任务停止前尚未处理此组合。", "candidates": []}
                    products = pair["candidates"]
                    statuses[pair["status"]] += 1
                    pairs.append({**base, "status": pair["status"], "candidate_count": len(products),
                                  "error_code": pair["error_code"], "message": pair["message"]})
                    for number, candidate in enumerate(products, 1):
                        record = {**base, "a_name": a["name"], "b_name": b["name"],
                                  "a_input_smiles": a["input_smiles"], "b_input_smiles": b["input_smiles"],
                                  "a_canonical_smiles": a["canonical_smiles"], "b_canonical_smiles": b["canonical_smiles"],
                                  "candidate_index": number, **candidate}
                        results.append(record)
                        if candidate_count % 50 == 0:
                            preview_offsets.append(preview_file.tell())
                        encoded = (json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n").encode()
                        budget.add(len(encoded))
                        preview_file.write(encoded)
                        candidate_count += 1
            if statuses["not_processed"] and not job["terminal_intent"]:
                raise RuntimeError("unprocessed pairs in a completed job")
            status = job["terminal_intent"] or ("completed_with_errors" if statuses["error"] or statuses["invalid_input"] else "completed")
            summary = {**job["summary"], "pair_statuses": dict(statuses), "candidate_count": candidate_count,
                       "processed_pairs": statuses["success"] + statuses["no_match"] + statuses["error"],
                       "pair_errors": statuses["error"],
                       "computed_unique_pairs": len(index.locations), "complete": status in {"completed", "completed_with_errors"},
                       "target_class": job["options"]["target_class"], "engine": job["engine"],
                       "limits": job["options"].get("limits", {}),
                       "input_files": snapshot.get("files", {}),
                       "candidate_definition": "SMiPoly 返回并通过跨表来源过滤的全部候选；保留原始行来源。",
                       "status": status, "error_code": job.get("error_code"), "message": job.get("message")}
        # CSVs are closed before hashing and packaging.
        csv_files = sorted(directory.glob("*.csv"))
        summary["csv_files"] = {path.name: {"sha256": digest(path), "size_bytes": path.stat().st_size} for path in csv_files}
        write_json(directory / "summary.json", summary)
        with ExitStack() as summary_stack:
            summary_writer = TableWriter(directory, "summary", ["key", "value"], workbook, budget, summary_stack)
            for key, value in summary.items():
                summary_writer.append({"key": key, "value": value})
        csv_files.append(directory / "summary.csv")
        workbook.save(directory / "results.xlsx")
        write_json(directory / "preview-index.json", {"offsets": preview_offsets, "count": candidate_count})
        with zipfile.ZipFile(directory / "results.zip", "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
            for path in [*csv_files, directory / "summary.json"]:
                archive.write(path, path.name)
        total_bytes = sum(path.stat().st_size for path in directory.iterdir() if path.is_file())
        if total_bytes + sum(chunk["artifact"]["size_bytes"] for chunk in chunks) > config.result_bytes:
            raise BatchError("导出文件总量超过存储限额。", "result_limit", 413)
        artifacts = {path.name: file_manifest(root, path, "text/csv; charset=utf-8") for path in csv_files}
        for name, media in [("results.zip", "application/zip"), ("results.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                            ("summary.json", "application/json"), ("preview.jsonl", "application/x-ndjson"), ("preview-index.json", "application/json")]:
            artifacts[name] = file_manifest(root, directory / name, media)
        return {"status": status, "summary": summary, "artifacts": artifacts}
    finally:
        index.close()
        workbook.close()
