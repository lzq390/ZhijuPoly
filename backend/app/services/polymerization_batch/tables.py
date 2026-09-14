from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.etree.ElementTree import ParseError

from app.utils.exceptions import InvalidSmilesError

from .chemistry import canonicalize
from .models import BatchError, BatchSettings, TableMapping


ALIASES = {
    "smiles_column": {"smiles", "canonical_smiles", "单体smiles", "单体 smiles", "结构smiles"},
    "id_column": {"id", "monomer_id", "编号", "单体编号", "cid"},
    "name_column": {"name", "monomer_name", "名称", "单体名称"},
}
EXCEL_CELL_CHARS = 32767
INPUT_HEADER_PREFIX = "input:"
INVALID_XML_TEXT = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def text_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def checked_workbook(path: Path, config: BatchSettings):
    from openpyxl import load_workbook
    try:
        with zipfile.ZipFile(path) as package:
            entries = package.infolist()
            if len(entries) > 2048 or sum(item.file_size for item in entries) > config.xlsx_uncompressed_bytes:
                raise BatchError("Excel 解压后内容超过限额。", "xlsx_limit", 413)
            if any(item.flag_bits & 1 for item in entries):
                raise BatchError("不支持加密 Excel 文件。")
        return load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except BatchError:
        raise
    except Exception as exc:
        raise BatchError("Excel 文件损坏或格式不受支持。", "invalid_xlsx") from exc


def read_table(path: Path, mapping: TableMapping, config: BatchSettings) -> dict:
    workbook = None
    try:
        if path.suffix == ".xlsx":
            workbook = checked_workbook(path, config)
            sheets = [sheet.title for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
            selected = mapping.sheet or (sheets[0] if sheets else None)
            if selected not in sheets:
                raise BatchError("请选择有效的可见工作表。", "invalid_sheet")
            sheet = workbook[selected]
            if sheet.max_column and sheet.max_column > config.max_columns:
                raise BatchError("Excel 列数超过限额。", "column_limit", 413)
            # Ignore unreliable worksheet dimensions; enforce actual limits below.
            sheet.reset_dimensions()
            iterator = ((number, [text_cell(cell.value) for cell in cells],
                         {index for index, cell in enumerate(cells) if cell.data_type == "f"})
                        for number, cells in enumerate(sheet.iter_rows(), 1))
        else:
            sheets, selected = [], None
            try:
                content = path.read_text(encoding=mapping.encoding)
            except UnicodeError as exc:
                raise BatchError("CSV 编码不匹配，请选择 UTF-8 或 GB18030 后重新预检。", "invalid_encoding") from exc
            csv.field_size_limit(config.max_cell_chars * 4)
            reader = csv.reader(io.StringIO(content, newline=""), strict=True)
            # A CSV record may span physical lines. row_number is its 1-based
            # table record number, including the header, just as in Excel.
            iterator = ((number, values, set()) for number, values in enumerate(reader, 1))
        headers = None
        rows, blank_rows = [], 0
        for number, values, formulas in iterator:
            if number > max(config.max_rows * 4, 20000):
                raise BatchError("表格物理记录数超过限额，请移除多余空行。", "row_limit", 413)
            if len(values) > config.max_columns or any(len(value) > config.max_cell_chars for value in values):
                raise BatchError("列数或单元格长度超过限额。", "table_limit", 413)
            if any(INVALID_XML_TEXT.search(value) for value in values):
                raise BatchError(f"第 {number} 行包含 Excel 无法保存的控制字符，请清理后重新上传。", "invalid_cell")
            if headers is None:
                headers = [value.strip() for value in values]
                if not headers or any(not value for value in headers) or len(set(headers)) != len(headers) or formulas:
                    raise BatchError("第 1 行必须包含非空且不重复的表头。", "invalid_header")
                if any(len(value) + len(INPUT_HEADER_PREFIX) > EXCEL_CELL_CHARS for value in headers):
                    raise BatchError("表头过长，无法完整保存至 Excel 输出，请缩短后重新上传。", "header_limit", 413)
                continue
            if not any(value.strip() for value in values):
                blank_rows += 1
                continue
            if len(values) > len(headers):
                raise BatchError(f"第 {number} 行的字段数超过表头。", "invalid_row")
            values += [""] * (len(headers) - len(values))
            rows.append({"row_number": number, "values": dict(zip(headers, values)),
                         "formula_columns": [headers[index] for index in formulas if index < len(headers)]})
            if len(rows) > config.max_rows:
                raise BatchError(f"每表最多 {config.max_rows} 条非空记录。", "row_limit", 413)
        if headers is None:
            raise BatchError("上传表为空。", "empty_table")
        return {"headers": headers, "sheets": sheets, "sheet": selected, "rows": rows,
                "blank_rows": blank_rows, "encoding": mapping.encoding}
    except csv.Error as exc:
        raise BatchError("CSV 格式错误或单元格超过限额。", "invalid_csv") from exc
    except (ParseError, ValueError, TypeError, KeyError, IndexError, EOFError, zipfile.BadZipFile) as exc:
        # read_only workbooks parse worksheet contents lazily during iteration.
        # Keep those failures in the same per-file import error contract as
        # errors raised by load_workbook itself.
        if workbook is not None:
            raise BatchError("Excel 工作表内容损坏或格式不受支持。", "invalid_xlsx") from exc
        raise
    finally:
        if workbook is not None:
            workbook.close()


def inspect_table(path: Path, mapping: TableMapping, config: BatchSettings) -> dict:
    table = read_table(path, mapping, config)
    defaults = mapping.model_dump()
    defaults["sheet"] = table["sheet"]
    for field, aliases in ALIASES.items():
        if defaults[field] is None:
            matches = [header for header in table["headers"] if header.lower() in aliases]
            defaults[field] = matches[0] if len(matches) == 1 else None
    return {"headers": table["headers"], "sheets": table["sheets"], "sheet": table["sheet"],
            "row_count": len(table["rows"]), "blank_rows": table["blank_rows"],
            "sample": [row["values"] for row in table["rows"][:20]], "mapping": defaults}


def validate_table(path: Path, role: str, mapping: TableMapping, config: BatchSettings) -> dict:
    table = read_table(path, mapping, config)
    for field in ("smiles_column", "id_column", "name_column"):
        column = getattr(mapping, field)
        if (field == "smiles_column" and not column) or (column and column not in table["headers"]):
            raise BatchError("请选择有效的 SMILES、编号和名称列。", "invalid_mapping")
    cache: dict[str, tuple[str | None, str]] = {}
    records = []
    for row in table["rows"]:
        raw = row["values"][mapping.smiles_column]
        error, canonical = "", None
        if mapping.smiles_column in row["formula_columns"]:
            error = "SMILES 单元格不能使用公式，请粘贴文本值。"
        elif not raw.strip():
            error = "SMILES 为空。"
        else:
            if raw not in cache:
                try:
                    cache[raw] = (canonicalize(raw), "")
                except (ValueError, InvalidSmilesError) as exc:
                    # Ordinary chemistry validation errors belong to this row.
                    cache[raw] = (None, str(exc)[:500])
            canonical, error = cache[raw]
        records.append({"source_key": f"{role}:{row['row_number']}", "row_number": row["row_number"],
                        "id": row["values"].get(mapping.id_column, "") if mapping.id_column else "",
                        "name": row["values"].get(mapping.name_column, "") if mapping.name_column else "",
                        "input_smiles": raw, "canonical_smiles": canonical,
                        "error_code": "invalid_smiles" if error else "", "message": error,
                        "values": row["values"]})
    valid = [row for row in records if row["canonical_smiles"]]
    unique = sorted({row["canonical_smiles"] for row in valid})
    return {"role": role, "headers": table["headers"], "mapping": mapping.model_dump(),
            "rows": records, "row_count": len(records), "valid_rows": len(valid),
            "unique_count": len(unique), "duplicate_rows": len(valid) - len(unique),
            "invalid_rows": len(records) - len(valid), "blank_rows": table["blank_rows"], "unique_smiles": unique}


def preview(files: dict, mappings: dict, root: Path, config: BatchSettings) -> dict:
    tables = {role: validate_table(root / files[role]["path"], role, TableMapping(**mappings[role]), config)
              for role in ("a", "b")}
    raw_pairs = tables["a"]["row_count"] * tables["b"]["row_count"]
    if raw_pairs > config.max_pairs:
        raise BatchError(f"原始组合数 {raw_pairs} 超过 {config.max_pairs} 对限额。", "pair_limit", 413)
    return {"tables": tables, "raw_pairs": raw_pairs,
            "valid_pairs": tables["a"]["valid_rows"] * tables["b"]["valid_rows"],
            "unique_pairs": tables["a"]["unique_count"] * tables["b"]["unique_count"],
            "can_submit": all(table["valid_rows"] > 0 for table in tables.values())}
