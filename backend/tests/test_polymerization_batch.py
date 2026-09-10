from __future__ import annotations

import csv
import io
import json
import zipfile
from dataclasses import replace
from datetime import timedelta
from pathlib import Path
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.config import Settings
from app.main import create_app
from app.models import MonomerPolymerizationRequest
from app.services.monomer_polymerization import run_monomer_polymerization
from app.services.polymerization_batch.chemistry import canonicalize, classify, engine_fingerprint, generate
from app.services.polymerization_batch.execution import classify_isolated, generate_isolated
from app.services.polymerization_batch.models import BatchError, BatchSettings, TableMapping
from app.services.polymerization_batch.service import BASE_PATH, BatchService
from app.services.polymerization_batch.storage import read_json, write_json
from app.services.polymerization_batch.tables import inspect_table, validate_table
from app.services.polymerization_batch.worker import BatchWorker
from app.services.deployment_control import count_active_postgres_jobs, enable_drain, disable_drain


DIAMINE = "Nc1ccc(N)cc1"
DIANHYDRIDE = "O=C1OC(=O)c2cc3c(cc21)C(=O)OC3=O"


@pytest.fixture
def batch(tmp_path, postgres_dsn):
    config = BatchSettings(enabled=True, storage_root=tmp_path / "batch", chunk_size=1)
    service = BatchService(postgres_dsn, config)
    app = create_app(Settings(app_postgres_dsn=postgres_dsn, model_enabled=False, retro_model_enabled=False))
    app.state.polymerization_batch = service
    worker = BatchWorker(postgres_dsn, config)
    worker.engine = engine_fingerprint()
    worker.heartbeat(force=True)
    with service.repository.connection() as conn:
        conn.execute("UPDATE governance.deployment_control SET drain_enabled=false,reason=NULL,release_sha=NULL,activated_by=NULL")
    return service, worker, TestClient(app)


def import_pair(client, a=None, b=None):
    a = a or f"id,name,SMILES\n001,二胺,{DIAMINE}\n001,重复二胺,{DIAMINE}\nbad,错误,not-smiles\n"
    b = b or f"id,SMILES\nB1,{DIANHYDRIDE}\nB2,CCO\n"
    response = client.post(BASE_PATH + "/imports", files={"file_a": ("../二胺.csv", a.encode("utf-8-sig"), "text/csv"), "file_b": ("二酐.csv", b.encode(), "text/csv")})
    assert response.status_code == 201, response.text
    uploaded = response.json()
    response = client.post(f"{BASE_PATH}/imports/{uploaded['import_id']}/preview", json={role: uploaded["tables"][role]["mapping"] for role in ("a", "b")})
    assert response.status_code == 200, response.text
    return response.json()


def submit(client, preview, key=None):
    return client.post(BASE_PATH + "/jobs", headers={"Idempotency-Key": key or uuid4().hex}, json={"import_id": preview["import_id"], "preview_revision": preview["preview_revision"], "target_class": "polyimide"})


def finish(worker, job_id):
    for _ in range(30):
        claimed = worker.repository.claim(worker.worker_id, worker.engine["fingerprint"])
        if claimed:
            worker.execute(*claimed)
        job = worker.repository.get_job(job_id)
        if job["status"] not in {"queued", "running", "cancelling"}:
            return job
    raise AssertionError("worker did not finish")


def test_real_batch_upload_compute_export_and_pagination(batch):
    service, worker, client = batch
    preview = import_pair(client)
    assert preview["statistics"]["raw_pairs"] == 6
    assert preview["statistics"]["valid_pairs"] == 4
    assert preview["statistics"]["unique_pairs"] == 2
    assert preview["input_error_count"] == 1
    key = uuid4().hex
    response = submit(client, preview, key)
    assert response.status_code == 202, response.text
    job_id = response.json()["job_id"]
    assert submit(client, preview, key).json()["job_id"] == job_id
    job = finish(worker, job_id)
    assert job["status"] == "completed_with_errors", job
    assert job["summary"]["candidate_count"] == 2
    assert job["summary"]["pair_statuses"] == {"success": 2, "no_match": 2, "invalid_input": 2}
    results = client.get(f"{BASE_PATH}/jobs/{job_id}/results?limit=1").json()
    assert results["total"] == 2 and results["next_offset"] == 1
    assert results["items"][0]["a_id"] == "001"
    assert results["items"][0]["a_input_smiles"] == DIAMINE
    assert results["items"][0]["engine_mon1_smiles"] != DIAMINE
    archive_response = client.get(f"{BASE_PATH}/jobs/{job_id}/artifacts/results.zip")
    assert archive_response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(archive_response.content)) as archive:
        rows = list(csv.DictReader(io.StringIO(archive.read("pairs.csv").decode("utf-8-sig"))))
        assert len(rows) == 6
        assert json.loads(archive.read("summary.json"))["complete"] is True
    excel_response = client.get(f"{BASE_PATH}/jobs/{job_id}/artifacts/results.xlsx")
    workbook = load_workbook(io.BytesIO(excel_response.content), read_only=True)
    assert len(list(workbook["results"].rows)) == 3
    assert {"inputs_a", "inputs_b", "pairs", "input_errors", "summary"}.issubset(workbook.sheetnames)
    workbook.close()
    assert not (service.root.parent / "二胺.csv").exists()


def test_recovery_fences_old_attempt_and_preserves_committed_chunks(batch):
    service, worker, client = batch
    preview = import_pair(client)
    job_id = submit(client, preview).json()["job_id"]
    first = worker.repository.claim(worker.worker_id, worker.engine["fingerprint"])
    worker.execute(*first)
    before = worker.repository.chunks(job_id, completed_only=True)
    abandoned = worker.repository.claim(worker.worker_id, worker.engine["fingerprint"])
    worker.repository.recover()
    assert not worker.repository.finish_unit(*abandoned, {"path": "untrusted", "size_bytes": 0}, 1)
    restarted = BatchWorker(service.repository.dsn, service.config)
    restarted.engine = worker.engine
    job = finish(restarted, job_id)
    after = worker.repository.chunks(job_id, completed_only=True)
    assert after[0]["artifact"] == before[0]["artifact"]
    assert job["summary"]["candidate_count"] == 2


def test_cancel_retains_complete_pairs_accounting(batch):
    _, worker, client = batch
    preview = import_pair(client)
    job_id = submit(client, preview).json()["job_id"]
    assert client.post(f"{BASE_PATH}/jobs/{job_id}/cancel").json()["status"] == "cancelling"
    assert client.post(f"{BASE_PATH}/jobs/{job_id}/cancel").status_code == 200
    job = finish(worker, job_id)
    assert job["status"] == "cancelled"
    assert job["summary"]["pair_statuses"] == {"not_processed": 4, "invalid_input": 2}
    assert job["summary"]["complete"] is False
    assert "results.zip" in job["artifacts"]


def test_stale_preview_empty_side_and_idempotency_conflict(batch):
    _, _, client = batch
    first = import_pair(client)
    key = uuid4().hex
    assert submit(client, first, key).status_code == 202
    second = import_pair(client)
    assert submit(client, second, key).status_code == 409
    replaced = client.post(f"{BASE_PATH}/imports/{first['import_id']}/preview", json={role: first["tables"][role]["mapping"] for role in ("a", "b")})
    assert replaced.status_code == 200
    assert submit(client, first).status_code == 409
    empty = import_pair(client, a="id,SMILES\nx,invalid\n")
    assert not empty["can_submit"]
    assert submit(client, empty).status_code == 422


def test_drain_blocks_claim_but_pending_jobs_do_not_block_deployment(batch):
    service, worker, client = batch
    preview = import_pair(client)
    submit(client, preview)
    with service.repository.connection() as conn:
        summary = count_active_postgres_jobs(conn)
        assert summary.active_jobs_schema_version == 3
        assert summary.counts["polymerization_batch"] == 0
        enable_drain(conn, reason="test", activated_by="test", release_sha="a" * 40)
    try:
        assert worker.repository.claim(worker.worker_id, worker.engine["fingerprint"]) is None
    finally:
        with service.repository.connection() as conn:
            disable_drain(conn, expected_activated_by="test", expected_release_sha="a" * 40)
    claim = worker.repository.claim(worker.worker_id, worker.engine["fingerprint"])
    with service.repository.connection() as conn:
        assert count_active_postgres_jobs(conn).counts["polymerization_batch"] == 1
    worker.execute(*claim)
    with service.repository.connection() as conn:
        assert count_active_postgres_jobs(conn).counts["polymerization_batch"] == 0


def test_retention_and_open_download_survive_unlink(batch):
    service, worker, client = batch
    job_id = submit(client, import_pair(client)).json()["job_id"]
    finish(worker, job_id)
    handle, _ = service.artifact(job_id, "results.zip")
    with service.repository.connection() as conn:
        conn.execute("UPDATE polymerization_batch.jobs SET expires_at=now()-interval '1 second' WHERE id=%s", (job_id,))
    worker.cleanup()
    with handle:
        assert handle.read(2) == b"PK"
    assert client.get(f"{BASE_PATH}/jobs/{job_id}/artifacts/results.zip").status_code == 410
    assert client.get(f"{BASE_PATH}/jobs/{job_id}").json()["status"] == "expired"


def test_xlsx_mapping_formulas_and_numeric_ids(tmp_path):
    path = tmp_path / "input.xlsx"
    workbook = Workbook()
    workbook.active.title = "ignore"
    workbook.active.append(["SMILES"])
    sheet = workbook.create_sheet("monomers")
    sheet.append(["编号", "结构", "备注"])
    sheet.append(["001", DIAMINE, "=1+1"])
    sheet.append([42, "=A1", "formula smiles"])
    sheet.append(["003", "", "missing"])
    workbook.save(path)
    config = BatchSettings(storage_root=tmp_path)
    mapping = TableMapping(sheet="monomers", smiles_column="结构", id_column="编号")
    table = validate_table(path, "a", mapping, config)
    assert table["valid_rows"] == 1 and table["invalid_rows"] == 2
    assert table["rows"][0]["id"] == "001" and table["rows"][1]["id"] == "42"
    assert "公式" in table["rows"][1]["message"]
    assert inspect_table(path, mapping, config)["sheets"] == ["ignore", "monomers"]
    with pytest.raises(BatchError):
        inspect_table(path, mapping, replace(config, xlsx_uncompressed_bytes=1))


def test_csv_encoding_headers_and_limits(tmp_path):
    path = tmp_path / "input.csv"
    path.write_bytes(f'编号,SMILES,名称\n001,{DIAMINE},"中文\n名称"\n\n'.encode("gb18030"))
    mapping = TableMapping(encoding="gb18030", smiles_column="SMILES", id_column="编号")
    table = validate_table(path, "a", mapping, BatchSettings())
    assert table["rows"][0]["row_number"] == 2 and table["blank_rows"] == 1
    with pytest.raises(BatchError, match="编码"):
        validate_table(path, "a", TableMapping(smiles_column="SMILES"), BatchSettings())
    path.write_text("SMILES,SMILES\nCC,CC\n")
    with pytest.raises(BatchError, match="表头"):
        inspect_table(path, TableMapping(), BatchSettings())


@pytest.mark.parametrize("target", ["polyimide", "polyamide", "polyester", "polyurethane", "polyether", "polyolefin", "polyoxazolidone", "all"])
def test_real_engine_cross_table_results_match_each_pair(target):
    a = [canonicalize(value) for value in [DIAMINE, "NCC(C)N", "OCCO", "C1CO1", "C=C", "O=C=NCCN=C=O"]]
    b = [canonicalize(value) for value in [DIANHYDRIDE, "O=C(O)CC(C)C(=O)O", "[C-]#[O+]", "OCCO", "C1CO1", DIAMINE]]
    cache = classify(sorted(set(a) | set(b)))
    assert len(cache["rows"]) == len(set(a) | set(b))
    generated = generate(a, b, cache, target)
    for pair in generated["pairs"]:
        baseline = run_monomer_polymerization(MonomerPolymerizationRequest(monomer_a_smiles=pair["a"], monomer_b_smiles=pair["b"], target_class=target, max_results=20))
        assert {(item["polymer_smiles"], item["polymer_class"], item["reaction_id"]) for item in pair["candidates"]} == {(item.polymer_smiles, item.polymer_class, item.reaction_id) for item in baseline.results}


def test_classification_and_generation_split_only_failing_units():
    def classify_call(request):
        if "bad" in request["smiles"]:
            raise BatchError("classification failed", "calculation_error")
        return {"rows": [{"source_key": item} for item in request["smiles"]], "errors": {}}
    cache = classify_isolated(["ok", "bad", "ok2"], classify_call)
    assert set(cache["errors"]) == {"bad"} and len(cache["rows"]) == 2
    def generate_call(request):
        if "bad" in request["a"]:
            raise BatchError("pair failed", "calculation_error")
        return {"pairs": [{"a": a, "b": b, "status": "no_match"} for a in request["a"] for b in request["b"]]}
    rows = list(generate_isolated(["ok", "bad"], ["b1", "b2"], generate_call))
    assert len(rows) == 4 and sum(row["status"] == "error" for row in rows) == 2


def test_export_has_no_twenty_result_cap_and_splits_excel_sheets(batch, monkeypatch):
    from app.services.polymerization_batch import exporting
    service, worker, client = batch
    preview = import_pair(client, a="id,SMILES\n" + "".join(f"{i:03d},{DIAMINE}\n" for i in range(25)),
                          b=f"id,SMILES\nB,{DIANHYDRIDE}\n")
    job_id = submit(client, preview).json()["job_id"]
    job = finish(worker, job_id)
    assert job["summary"]["candidate_count"] == 25
    assert job["summary"]["computed_unique_pairs"] == 1
    assert len(client.get(f"{BASE_PATH}/jobs/{job_id}/results").json()["items"]) == 25
    monkeypatch.setattr(exporting, "EXCEL_ROWS", 10)
    snapshot = read_json(service.root / job["options"]["snapshot_path"])
    output = service.root / "split-export"
    exporting.export_job(service.root, job, snapshot, worker.repository.chunks(job_id, completed_only=True), output, service.config)
    workbook = load_workbook(output / "results.xlsx", read_only=True)
    assert [name for name in workbook.sheetnames if name.startswith("results")] == ["results", "results_2", "results_3"]
    assert sum(sum(1 for _ in workbook[name].rows) - 1 for name in workbook.sheetnames if name.startswith("results")) == 25
    summary = {row[0]: row[1] if len(row) > 1 else None
               for name in workbook.sheetnames if name.startswith("summary") for row in workbook[name].values}
    assert "results.csv" in json.loads(summary["csv_files"])
    workbook.close()


def test_failed_mapping_invalidates_old_preview_and_old_request_cannot_publish(batch):
    service, _, client = batch
    preview = import_pair(client)
    # An unmappable new request must invalidate the previously valid revision.
    response = client.post(f"{BASE_PATH}/imports/{preview['import_id']}/preview", json={"a": {"smiles_column": "absent"}})
    assert response.status_code in {200, 422}
    assert submit(client, preview).status_code == 409
    old, new = uuid4().hex, uuid4().hex
    service.repository.begin_preview(preview["import_id"], old)
    service.repository.begin_preview(preview["import_id"], new)
    with pytest.raises(BatchError, match="更新的预检"):
        service.repository.save_preview(preview["import_id"], old, {})


def test_queue_capacity_preserves_idempotent_retry_when_worker_unavailable(batch):
    service, _, client = batch
    service.repository.config = replace(service.config, queue_capacity=1)
    preview = import_pair(client)
    key = uuid4().hex
    first = submit(client, preview, key).json()
    assert submit(client, preview).status_code == 202
    assert submit(client, preview).status_code == 429
    with service.repository.connection() as conn:
        conn.execute("UPDATE polymerization_batch.worker_status SET available=false")
    assert submit(client, preview, key).json()["job_id"] == first["job_id"]


def test_changed_engine_fails_without_mixing_committed_results(batch):
    service, worker, client = batch
    job_id = submit(client, import_pair(client)).json()["job_id"]
    worker.execute(*worker.repository.claim(worker.worker_id, worker.engine["fingerprint"]))
    committed = worker.repository.chunks(job_id, completed_only=True)
    assert worker.repository.claim(worker.worker_id, "different-version") is None
    job = service.repository.get_job(job_id)
    assert job["status"] == "failed" and job["error_code"] == "engine_version_mismatch"
    assert worker.repository.chunks(job_id, completed_only=True) == committed


def test_final_export_publish_failure_retries_without_duplicate_rows(batch, monkeypatch):
    service, worker, client = batch
    job_id = submit(client, import_pair(client)).json()["job_id"]
    while True:
        job, chunk = worker.repository.claim(worker.worker_id, worker.engine["fingerprint"])
        if chunk["kind"] == "export":
            break
        worker.execute(job, chunk)
    original = worker.repository.finish_unit
    monkeypatch.setattr(worker.repository, "finish_unit", lambda *args: False)
    worker.execute(job, chunk)
    assert service.repository.get_job(job_id)["artifacts"] == {}
    assert not (service.root / "jobs" / job_id / "attempts" / job["execution_token"]).exists()
    monkeypatch.setattr(worker.repository, "finish_unit", original)
    worker.repository.recover()
    completed = finish(worker, job_id)
    assert completed["summary"]["candidate_count"] == 2
    exported = next(item for item in worker.repository.chunks(job_id) if item["kind"] == "export")
    assert exported["attempt"] == 2 and exported["lease_expires_at"] is None


def test_system_errors_do_not_repeatedly_bisect():
    calls = []
    def fail(request):
        calls.append(request)
        raise BatchError("missing rules", "engine_unavailable", 503)
    with pytest.raises(BatchError):
        classify_isolated(["a", "b", "c"], fail)
    assert len(calls) == 1
    calls.clear()
    with pytest.raises(BatchError):
        list(generate_isolated(["a", "b"], ["c", "d"], fail))
    assert len(calls) == 1


@pytest.mark.parametrize("cancel", [False, True])
def test_timeout_and_cancellation_reap_the_child_process(tmp_path, monkeypatch, cancel):
    import subprocess
    import sys
    from app.services.polymerization_batch import execution
    original = subprocess.Popen
    children = []
    def start(arguments, **kwargs):
        child = original([sys.executable, "-c", "import time; time.sleep(30)"], **kwargs)
        children.append(child)
        return child
    monkeypatch.setattr(execution.subprocess, "Popen", start)
    def guard():
        if cancel:
            raise execution.ExecutionStopped("cancelled")
    expected = execution.ExecutionStopped if cancel else BatchError
    with pytest.raises(expected):
        execution.run_isolated({}, BatchSettings(), tmp_path, timeout=0.01, guard=guard)
    assert children[0].poll() is not None
    assert list(tmp_path.iterdir()) == []


def test_resource_limit_never_reports_truncated_success(batch, monkeypatch):
    from app.services.polymerization_batch import worker as worker_module
    service, worker, client = batch
    job_id = submit(client, import_pair(client)).json()["job_id"]
    original = worker_module.run_isolated
    def limited(request, *args, **kwargs):
        if request["action"] == "export":
            raise BatchError("result limit", "result_limit", 413)
        return original(request, *args, **kwargs)
    monkeypatch.setattr(worker_module, "run_isolated", limited)
    job = finish(worker, job_id)
    assert job["status"] == "failed" and job["error_code"] == "result_limit"
    assert not job["artifacts"]
    assert service.repository.chunks(job_id, completed_only=True)


def test_drain_freezes_worker_heartbeat_and_cleanup(batch):
    service, worker, client = batch
    preview = import_pair(client)
    with service.repository.connection() as conn:
        previous = conn.execute("SELECT * FROM polymerization_batch.worker_status").fetchone()
        enable_drain(conn, reason="test", activated_by="test", release_sha="a" * 40)
    try:
        worker.heartbeat(force=True)
        worker.cleanup()
        with service.repository.connection() as conn:
            assert conn.execute("SELECT * FROM polymerization_batch.worker_status").fetchone() == previous
        assert service.repository.get_import(preview["import_id"])
    finally:
        with service.repository.connection() as conn:
            disable_drain(conn, expected_activated_by="test", expected_release_sha="a" * 40)


def test_normal_export_obeys_overall_execution_budget(batch):
    _, worker, client = batch
    job_id = submit(client, import_pair(client)).json()["job_id"]
    while True:
        job, chunk = worker.repository.claim(worker.worker_id, worker.engine["fingerprint"])
        if chunk["kind"] == "export":
            break
        worker.execute(job, chunk)
    worker.config = replace(worker.config, job_seconds=1)
    worker.execute(job, chunk)
    stopped = worker.repository.get_job(job_id)
    assert stopped["status"] == "failed" and stopped["error_code"] == "job_timeout"
    assert not stopped["artifacts"]
