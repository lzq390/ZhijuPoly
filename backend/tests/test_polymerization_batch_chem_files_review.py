from __future__ import annotations

import csv
import io
import zipfile
from contextlib import ExitStack

import pytest
from openpyxl import Workbook, load_workbook

from app.services.polymerization_batch.exporting import OutputBudget, TableWriter, export_job
from app.services.polymerization_batch.models import BatchError, BatchSettings, TableMapping
from app.services.polymerization_batch.storage import file_manifest, write_json
from app.services.polymerization_batch.tables import EXCEL_CELL_CHARS, INPUT_HEADER_PREFIX, read_table


def test_cancelled_export_preserves_known_classification_errors(tmp_path):
    def row(role, number, smiles):
        return {"source_key": f"{role}:{number}", "row_number": number, "id": f"{number:03d}",
                "name": "", "input_smiles": smiles or "bad", "canonical_smiles": smiles,
                "error_code": "" if smiles else "invalid_smiles", "message": "" if smiles else "bad input",
                "values": {"SMILES": smiles or "bad"}}

    snapshot = {"tables": {
        "a": {"headers": ["SMILES"], "rows": [row("a", 2, "CC"), row("a", 3, "CCC"), row("a", 4, None)]},
        "b": {"headers": ["SMILES"], "rows": [row("b", 2, "NN"), row("b", 3, "OO")]},
    }}
    cache = tmp_path / "classification.json"
    write_json(cache, {"rows": [], "errors": {"CC": "classification failed"}})
    chunks = [{"kind": "classify", "artifact": file_manifest(tmp_path, cache, "application/json")}]
    result = export_job(tmp_path, {
        "terminal_intent": "cancelled", "summary": {"raw_pairs": 6, "pair_errors": 0},
        "options": {"target_class": "polyamide"}, "engine": {},
    }, snapshot, chunks, tmp_path / "export", BatchSettings())

    assert result["status"] == "cancelled"
    assert result["summary"]["complete"] is False
    assert result["summary"]["pair_statuses"] == {"error": 2, "not_processed": 2, "invalid_input": 2}
    assert result["summary"]["pair_errors"] == 2
    assert sum(result["summary"]["pair_statuses"].values()) == 6
    with zipfile.ZipFile(tmp_path / "export" / "results.zip") as archive:
        pairs = list(csv.DictReader(io.StringIO(archive.read("pairs.csv").decode("utf-8-sig"))))
        assert [pair["pair_id"] for pair in pairs] == ["a2:b2", "a2:b3", "a3:b2", "a3:b3", "a4:b2", "a4:b3"]
        assert all(pair["error_code"] == "classification_error" for pair in pairs[:2])
        errors = list(csv.DictReader(io.StringIO(archive.read("input_errors.csv").decode("utf-8-sig"))))
        assert errors[0]["source_key"] == "a:2" and errors[0]["message"] == "classification failed"


@pytest.mark.parametrize("bad_cell", [b'<c r="A2" t="n"><v>not-a-number</v></c>', b'<c r="A2"><v>1</v></wrong>'])
def test_lazy_xlsx_worksheet_errors_use_input_error_contract(tmp_path, bad_cell):
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.append(["SMILES"])
    workbook.active.append(["CC"])
    workbook.save(source)
    broken = tmp_path / "broken.xlsx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(broken, "w") as changed:
        for item in original.infolist():
            value = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                start, end = value.index(b'<row r="2">'), value.index(b'</row>', value.index(b'<row r="2">'))
                value = value[:start] + b'<row r="2">' + bad_cell + value[end:]
            changed.writestr(item, value)
    with pytest.raises(BatchError) as caught:
        read_table(broken, TableMapping(), BatchSettings())
    assert caught.value.code == "invalid_xlsx"


@pytest.mark.parametrize("control", ["\x00", "\x01", "\x0b", "\ufffe"])
def test_csv_rejects_unexportable_control_characters_before_computation(tmp_path, control):
    path = tmp_path / "input.csv"
    path.write_text(f"SMILES,name\nCC,example{control}name\n")
    with pytest.raises(BatchError) as caught:
        read_table(path, TableMapping(), BatchSettings())
    assert caught.value.code == "invalid_cell"


def test_import_header_limit_accounts_for_output_provenance_prefix(tmp_path):
    path = tmp_path / "input.csv"
    accepted = "x" * (EXCEL_CELL_CHARS - len(INPUT_HEADER_PREFIX))
    path.write_text(f"SMILES,{accepted}\nCC,text\n")
    assert read_table(path, TableMapping(), BatchSettings())["headers"] == ["SMILES", accepted]
    path.write_text(f"SMILES,{accepted}x\nCC,text\n")
    with pytest.raises(BatchError) as caught:
        read_table(path, TableMapping(), BatchSettings())
    assert caught.value.code == "header_limit"


def test_export_rejects_overlong_header_instead_of_silent_xlsx_truncation(tmp_path):
    workbook = Workbook(write_only=True)
    with ExitStack() as stack, pytest.raises(BatchError) as caught:
        TableWriter(tmp_path, "inputs_a", ["x" * (EXCEL_CELL_CHARS + 1)], workbook, OutputBudget(1_000_000), stack)
    assert caught.value.code == "excel_cell_limit"
    workbook.close()


def test_excel_export_preserves_safe_text_and_leading_zero_ids(tmp_path):
    workbook = Workbook(write_only=True)
    with ExitStack() as stack:
        writer = TableWriter(tmp_path, "inputs_a", ["id", "name", "note"], workbook, OutputBudget(1_000_000), stack)
        writer.append({"id": "001", "name": "=1+1", "note": "中文\n换行\t内容"})
    path = tmp_path / "result.xlsx"
    workbook.save(path)
    read = load_workbook(path, read_only=True, data_only=False)
    values = list(read["inputs_a"].values)
    assert values[1] == ("001", "'=1+1", "中文\n换行\t内容")
    assert all(cell.data_type == "s" for cell in list(read["inputs_a"].rows)[1])
    read.close()
